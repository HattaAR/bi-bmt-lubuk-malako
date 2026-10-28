"""Bounded Excel reader; XLSX cells adapt to the existing xlrd validation rules."""
from datetime import date,datetime
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile,BadZipFile
import xlrd

MAX_UPLOAD_BYTES=50*1024*1024
MAX_CELLS=2_000_000


def check_container(name,contents):
    if not contents or len(contents)>MAX_UPLOAD_BYTES:
        raise ValueError('Ukuran berkas harus lebih dari nol dan maksimal 50 MB.')
    suffix=name.rsplit('.',1)[-1].lower()
    if suffix=='xls':
        if not contents.startswith(bytes.fromhex('D0CF11E0A1B11AE1')):
            raise ValueError('Isi berkas bukan workbook XLS yang didukung.')
    elif suffix=='xlsx':
        try:
            with ZipFile(BytesIO(contents)) as archive:
                items=archive.infolist();names={x.filename for x in items}
                if (len(items)>10000 or sum(x.file_size for x in items)>200*1024*1024
                    or any(x.flag_bits&1 or x.file_size>max(1,x.compress_size)*1000 for x in items)
                    or '[Content_Types].xml' not in names or 'xl/workbook.xml' not in names
                    or any('vba' in n.lower() or 'externallinks/' in n.lower() or 'embeddings/' in n.lower() for n in names)):
                    raise ValueError('Workbook terlalu besar atau memuat konten aktif/tautan eksternal yang tidak didukung.')
        except (BadZipFile,KeyError):
            raise ValueError('Isi berkas bukan workbook XLSX yang valid.') from None
    else:raise ValueError('Hanya .xls dan .xlsx yang didukung.')


def open_book(path,contents=None):
    contents=contents if contents is not None else path.read_bytes()
    check_container(path.name,contents)
    if path.suffix.lower()=='.xls':return xlrd.open_workbook(str(path),file_contents=contents)
    from openpyxl import load_workbook
    from openpyxl.utils.datetime import to_excel,CALENDAR_MAC_1904
    book=load_workbook(BytesIO(contents),read_only=True,data_only=False,keep_links=False)
    epoch=book.epoch
    class Sheet:
        def __init__(self,sheet):
            if sheet.max_row*sheet.max_column>MAX_CELLS:raise ValueError('Workbook melampaui batas sel pemeriksaan.')
            self.rows=[]
            for raw in sheet.iter_rows():
                cells=[]
                for cell in raw:
                    value=cell.value
                    if cell.data_type=='f':raise ValueError('Formula Excel tidak didukung; gunakan ekspor nilai asli.')
                    if value is None:kind,value=xlrd.XL_CELL_EMPTY,''
                    elif isinstance(value,(date,datetime)):kind,value=xlrd.XL_CELL_DATE,to_excel(value,epoch)
                    elif isinstance(value,bool):kind=xlrd.XL_CELL_BOOLEAN
                    elif cell.data_type=='e':kind=xlrd.XL_CELL_ERROR
                    elif isinstance(value,(int,float)):kind=xlrd.XL_CELL_NUMBER
                    else:kind=xlrd.XL_CELL_TEXT
                    cells.append(xlrd.sheet.Cell(kind,value))
                self.rows.append(cells)
            self.nrows=len(self.rows)
        def row_values(self,index):return [c.value for c in self.rows[index]]
        def row(self,index):return self.rows[index]
    names=book.sheetnames
    # Materialize only the explicitly selected sheet, never formulas or links.
    class Adapter:
        datemode=1 if epoch==CALENDAR_MAC_1904 else 0
        def sheet_names(self):return names
        def release_resources(self):book.close()
        def sheet_by_name(self,name):
            try:return Sheet(book[name])
            finally:book.close()
    return Adapter()
